from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output' / 'vla'
TMP = ROOT / 'tmp' / 'vla'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.8)
sec.left_margin, sec.right_margin = Cm(2.1), Cm(2.1)
sec.footer_distance = Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name]
    s.font.name='Microsoft YaHei'
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    s.font.color.rgb=RGBColor(0,0,0)
    s.font.italic=False
    s.paragraph_format.space_after=Pt(5)
    pr=s._element.find(qn('w:pPr'))
    if pr is not None:
        for border in list(pr.findall(qn('w:pBdr'))): pr.remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.17
doc.styles['Title'].font.size=Pt(29)
doc.styles['Subtitle'].font.size=Pt(14)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12.5)
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.space_before=Pt(9)
doc.core_properties.title='VLA 机器人模型入门讲义'
doc.core_properties.subject='视觉语言动作模型与机器人控制'
doc.core_properties.author='Codex'
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('VLA 入门讲义  |  '); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def p(text, bold=False):
    x=doc.add_paragraph(); r=x.add_run(text); r.bold=bold; return x
def h(text): doc.add_heading(text,level=2)
def page(num,title):
    title_p=doc.add_heading(f'{num:02d}  {title}',level=1)
    title_p.paragraph_format.page_break_before=True
def code(text):
    x=doc.add_paragraph()
    x.paragraph_format.line_spacing=1.05
    for i,line in enumerate(text.splitlines()):
        r=x.add_run(('\n' if i else '')+line)
        r.font.name='Consolas'; r.font.size=Pt(9)
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,(label,w) in enumerate(zip(headers,widths)):
        t.columns[i].width=Cm(w); t.rows[0].cells[i].width=Cm(w);t.rows[0].cells[i].text=label
    for row in rows:
        cells=t.add_row().cells
        for i,(s,w) in enumerate(zip(row,widths)): cells[i].text=s;cells[i].width=Cm(w)
    for j,row in enumerate(t.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            fill='DFE9F2' if j==0 else ('F5F7F9' if j%2==0 else 'FFFFFF')
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),fill);pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.line_spacing=1.12
                for r in para.runs:r.font.size=Pt(9);r.bold=(j==0)
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def diagram(name,kind):
    im=Image.new('RGB',(1500,560),'white');d=ImageDraw.Draw(im)
    f=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',31)
    small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',25)
    def box(rect,title,sub):
        d.rounded_rectangle(rect,16,fill='#EDF3F8',outline='#71879D',width=2)
        x0,y0,x1,y1=rect
        d.text(((x0+x1)/2,y0+32),title,font=f,fill='#152E46',anchor='mt')
        for j,s in enumerate(sub):d.text(((x0+x1)/2,y0+91+j*39),s,font=small,fill='#263B4D',anchor='mt')
    def arrow(a,b):
        d.line([a,b],fill='#304B66',width=5)
        x,y=b;d.polygon([(x,y),(x-17,y-10),(x-17,y+10)],fill='#304B66')
    if kind=='loop':
        boxes=[(30,60,330,250),(410,60,710,250),(790,60,1090,250),(1170,60,1470,250)]
        texts=[('观察与指令',['相机  状态  文本']),('VLA 策略',['预测动作或动作块']),('执行与控制',['适配  限幅  跟踪']),('机器人与环境',['产生新的观察'])]
        for rect,(title,sub) in zip(boxes,texts):box(rect,title,sub)
        for a,b in zip(boxes,boxes[1:]):arrow((a[2]+8,155),(b[0]-8,155))
        d.line([(1320,265),(1320,405),(180,405),(180,270)],fill='#304B66',width=5)
        d.polygon([(180,270),(170,287),(190,287)],fill='#304B66')
        d.text((750,436),'执行后重新观察  闭环修正偏差',font=f,fill='#152E46',anchor='mt')
    elif kind=='layers':
        box((30,100,440,325),'任务与语义',['VLA 或视觉语言系统','选择目标  生成参考'])
        box((545,100,955,325),'运动与平衡控制',['状态估计  MPC 或 PID','计算可执行控制量'])
        box((1060,100,1470,325),'驱动与硬件',['电机驱动  编码器  IMU','执行并反馈状态'])
        arrow((455,210),(530,210));arrow((970,210),(1045,210))
        d.text((750,405),'工程设计示意  控制接口与更新周期需要实测',font=f,fill='#152E46',anchor='mt')
    path=TMP/(name+'.png');im.save(path);doc.add_picture(str(path),width=Cm(16.7))

doc.add_paragraph('VLA 机器人模型入门讲义','Title')
doc.add_paragraph('视觉  语言  动作\n从基础概念到机器人控制实践','Subtitle')
p('面向刚开始学习机器人控制和人工智能的读者',True)
p('你已经在 Ubuntu 中运行过 TinyMPC。这份讲义以此为起点，解释机器人如何把相机画面和语言指令转成动作，以及 VLA 怎样与状态估计、轨迹规划和 MPC 配合。读完后，你应能画出一个完整的机器人系统，读懂代表论文的输入输出，并设计一个可检验的入门实验。')
p('VLA 是一类模型与学习方法的名称。它并不规定唯一网络结构，也不自动包含完整的机器人硬件、驱动、控制和安全系统。学习时，始终追问三件事：模型看到了什么，输出了什么，谁负责执行。')
diagram('overview','loop')
p('阅读方法',True)
p('第一次先读第 1 至 5 章，再读第 11 章的底盘例子。第二次补上训练、数据和代表模型，最后完成练习。公式用于解释变量关系，可以先读紧随其后的文字。每章末尾都给出检查理解的问题，答案集中在第 15 章。')
p('版本日期  2026 年 10 月 14 日\n代表模型按 2022 至 2025 年的关键方法选取，重点是原理与差异；本讲义不是最新模型排行榜。文中 [编号] 对应末尾可点击的一手资料。')

page(1,'学习地图与机器人基础')
table(['阶段','先解决的问题','对应章节'],[
('建立框架','观察 状态 动作 控制分别是什么','1 至 5'),
('理解学习','动作如何表示 数据怎样变成策略','6 至 8'),
('认识研究','模型有哪些路线 如何评估能力','9 至 10'),
('联系实践','与 MPC 配合 从小实验开始','11 至 15')],[2.6,10.9,3.3])
h('观察不等于完整状态')
p('状态是描述机器人和环境所需的信息，例如轮速、机身倾角、机械臂关节角和物体位置。观察是传感器实际给你的信息：相机图像、IMU 读数、编码器数据等。相机可能看不到被遮挡的杯子，IMU 也存在噪声，所以机器人通常只能获得状态的一部分。')
p('状态估计把多个传感器的读数与模型结合，推断更可靠的状态。对于平衡底盘，视觉能帮助发现目标，IMU 和编码器则更直接地反映机身姿态和运动。两种信息通常服务于不同的更新周期。')
h('策略把观察变成动作')
p('策略可以写成 aₜ = π(oₜ, l)。oₜ 是时刻 t 的观察，l 是语言指令，aₜ 是动作。这里的 π 是策略函数的符号，不特指 Physical Intelligence 的 π 系列模型。传统策略可以是手工规则；学习策略的参数来自训练数据。')
p('控制器负责使机器人跟踪目标。例如，上层给出“向前走”的速度参考，下层根据倾角和轮速计算电机控制量。执行动作后再次测量并修正，称为闭环；提前给出一串动作后不再看反馈，称为开环。')
h('用搬杯子区分不同任务')
p('“拿红杯子”是任务目标；找到杯子是感知；选择抓取姿态与路径是规划；跟踪关节轨迹是控制。VLA 可学习其中多个环节，但这些功能是否由同一个模型完成，取决于系统设计。')
p('自测 1  摄像头识别了前方目标，为什么平衡底盘仍然需要 IMU？')

page(2,'读懂 VLA 所需的人工智能概念')
h('参数 训练与推理')
p('神经网络是带有大量可调参数的函数。训练时，系统用样本计算误差，并修改参数，使输出更接近期望值；推理时，参数通常固定，模型对新输入给出预测。一次训练包含许多参数更新，一次推理则是完成一轮预测。')
p('损失函数衡量预测与监督信号之间的差异。动作回归可以比较预测动作与示教动作的数值，动作分类可以比较 token 的概率。训练损失降低表示模型更适配训练目标，仍需要独立的闭环实验验证机器人任务是否完成。')
h('Token 和 Transformer')
p('Token 是模型处理序列的基本单元。文字可拆成文本 token；图像可切成小块再变成向量；连续动作也可编码成离散 token。不同 token 的物理含义不同，不能把动作编号直接当成电机命令。')
p('Transformer 通过注意力机制，让序列里的表示根据其他表示调整自身。直观上，指令里的“红色”可以影响模型对图像不同区域的处理。注意力是计算机制，并不证明模型具有人的理解能力。原始 Transformer 论文提出以注意力为核心的架构。[1]')
h('预训练与微调')
p('预训练先用较广的数据学习可迁移的表示；微调再用目标机器人或任务的数据调整模型。预训练视觉语言模型能提供物体、语言和场景知识，但物理动作接口、接触和反馈仍需要机器人数据与执行系统。')
p('LoRA 是一种参数高效微调方式：冻结原有的大部分权重，训练较小的低秩更新矩阵。它能减少需要训练的参数和部分显存开销，但模型前向计算、数据质量和部署验证仍然不可省略。[2]')
h('模仿学习与强化学习')
p('模仿学习从示范中学习“这种情况下怎样动作”；行为克隆是其中最直接的一类。强化学习利用奖励和交互改进策略。VLA 可以采用模仿学习，也可以进一步结合强化学习；这两个词描述学习方式，VLA 描述视觉、语言与动作的关联。')
p('自测 2  下载模型后对图像预测动作，是训练还是推理？仅靠反复改指令能否代替机器人微调数据？')

page(3,'从视觉语言模型到视觉语言动作模型')
table(['名称','典型输入','典型输出'],[
('LLM 语言模型','文本','文本或结构化内容'),
('VLM 视觉语言模型','图像与文本','描述 回答 或其他表征'),
('VLA 视觉语言动作模型','图像 指令 可选状态与历史','机器人动作或动作序列')],[4.4,5.5,6.9])
h('把任务描述落实到可执行动作')
p('VLM 可能回答“红杯子在左边”。VLA 则需要根据当前场景与机器人接口，预测向哪里移动、如何改变姿态、何时闭合夹爪。动作必须与机器人能力、坐标系和控制周期对应。即使识别目标正确，动作也可能因为碰撞、夹爪偏差或时延而失败。')
p('常见 VLA 会继承预训练 VLM 的表示，再通过动作头或动作 token 学习控制。也有结构不同的视觉语言条件策略。“端到端”通常表示从观察到动作的映射通过学习完成；它并不意味着系统完全没有驱动、标定和底层伺服控制。')
h('语义知识与动作经验')
p('互联网数据有助于学习“杯子”“红色”“右边”等概念。机器人示教让模型学习这些概念如何对应实际运动。物体名称相同，不同机械臂的关节数、夹爪和相机位置仍可能不同，动作经验能否迁移需要验证。')
h('关于自主性的三个问题')
p('第一，模型能否选择正确目标？第二，能否执行到位？第三，失败后能否发现并恢复？同一系统可能在第一项表现好，而在后两项仍受限。因此，展示机器人成功拿起一次物体，不足以证明它能在所有场景里稳定执行。')
p('语言理解也要实测。同一张图，把指令从“拿红杯子”改成“拿蓝杯子”，动作是否相应变化？若没有变化，模型可能主要依赖视觉捷径或数据偏差。')
p('自测 3  能准确描述杯子位置的 VLM，是否已经是一个能抓杯子的 VLA？你还需要补哪些接口？')

page(4,'一次 VLA 推理如何运行')
diagram('closed_loop','loop')
h('输入到动作的路径')
p('相机提供图像，文本描述任务，状态输入提供关节或底盘信息。系统先按训练约定完成裁剪、缩放、颜色处理和状态归一化，再把这些输入交给模型。视觉编码器将图像转换为特征，语言模块处理指令，跨模态计算使它们相互关联。')
p('动作输出可能是一组离散 token，也可能由连续动作头产生。解码和反归一化将模型数值恢复为动作接口所需的单位。随后执行层完成时间调度、限幅和控制器调用。环境发生变化后，系统重新采集观察，再预测动作。')
h('一段动作与执行周期')
p('有些模型每次输出一个动作，有些输出未来 H 步的动作块。动作块在不同实现中代表位置目标、增量或其他控制量。执行全部 H 步期间，反馈可能较少；也可以只执行前 k 步，随后重新观察并预测。H 是预测长度，k 是实际执行长度，两者不必相同。')
p('长动作块可以减少模型调用次数，也可能让突然发生的环境变化更晚进入决策。采用短执行窗口、异步推理或动作衔接时，还需处理旧动作与新动作的时间对齐。')
h('不要混淆两种频率')
p('动作频率表示指令序列的时间间隔；模型调用频率表示多久重新运行一次模型。一个模型可能一次预测多个动作，然后由执行器按固定周期逐个发送。因而“生成 50 Hz 的动作”不自动等于“每 20 ms 完成一次全模型推理”。')
p('自测 4  若 H=20，动作间隔为 0.05 秒，则整块覆盖多久？如果每次只执行前 4 步，忽略推理时延后多久更新一次预测？')

page(5,'动作表示与坐标系')
h('动作向量里的数到底是什么')
p('教学示例：a = [Δx, Δy, Δz, Δroll, Δpitch, Δyaw, g]。前三个数可以表示末端位移增量，中间三个表示姿态增量，最后一个表示夹爪命令。这只是一个常见形式的示意，实际模型的顺序、单位、旋转定义和夹爪语义必须查配置。')
table(['表示','含义','部署时需确认'],[
('关节位置','各关节目标角度','角度单位 关节顺序 零点'),
('末端位姿或增量','工具的位置与旋转','基坐标或工具坐标 旋转表示'),
('速度','关节或底盘运动速度','单位 更新时间 执行模式'),
('力矩或力','驱动或接触控制量','硬件接口 动力学与边界')],[3.4,5.4,8.0])
h('同一个数字会产生不同动作')
p('沿工具坐标系 x 轴前进，与沿机器人基坐标系 x 轴前进，方向可能不同。旋转还可能使用欧拉角、旋转向量或四元数。四元数通常有四个分量；把四元数当作三个欧拉角使用，会改变动作意义。')
p('绝对位置与增量也不能混用。“0.02”可能是目标 x 坐标 0.02 米，也可能是从当前位置再移动 0.02 米。增量是否乘控制周期，取决于输出代表位移还是速度。')
h('归一化的作用')
p('为了让不同维度的数值尺度更可比，训练时常对动作做标准化或范围缩放。推理时必须使用匹配的数据统计量恢复动作。不要拿另一个机器人数据集的统计量直接解码；夹爪等离散或特殊维度也可能采用独立处理。')
p('假设教学数据采用 z=(a-μ)/σ。当 μ=0、σ=0.02 米、模型预测 z=0.5 时，原始位移 a=0.01 米。只有确实采用这种标准化的维度才能这样计算。')
p('自测 5  模型输出的 0.5 能否直接解释为移动 0.5 米？先检查哪几项？')

page(6,'离散动作与连续动作的生成')
h('把动作当作 token 预测')
p('离散方法将连续值划分为若干区间或采用其他编码，再预测动作 token。简单示意：把 [-1,1] 分成 256 个区间，每个维度用一个编号表示。编号再解码为区间代表值。编码方案越粗，能表达的精细程度越受限制；更复杂的 token 化可以压缩动作序列。')
p('自回归生成按顺序预测 token，每一步依赖前面已经生成的内容。它容易与语言模型的序列预测机制衔接，但一个动作块所需 token 数及其生成方式会影响延迟。RT-2 和原始 OpenVLA 是理解离散动作路线的代表。[4][5]')
h('连续回归与多种合理动作')
p('连续回归直接预测数值。若到目标有左右两条合理路径，训练数据各包含一种，把两条路径简单平均可能得到穿过障碍的路线。这是动作分布具有多种模式时会遇到的问题，不代表每一种回归策略必然失败。')
h('扩散策略与流匹配')
p('扩散式策略把动作序列作为生成对象，从噪声开始，通过多步条件去噪得到动作。Diffusion Policy 是这条路线的重要基础工作，其方法包含条件动作序列生成和滚动执行。[8]')
p('流匹配学习一个随时间变化的向量场，指引样本从初始噪声分布走向数据分布。生成时用数值积分得到动作。这里的生成时间是模型内部的变量，不等于机器人动作发生的真实时间。流匹配与扩散相关，但训练目标和构造不完全相同。[9]')
p('π0 将预训练视觉语言表示与流匹配动作生成结合。[6] 学习这一部分时，先掌握“根据观察生成一段连续动作”，再读向量场和积分公式。输出平滑程度、推理延迟和任务成功率仍要分别测量。')
h('动作块不等于 MPC')
p('VLA 根据学到的策略预测动作块；MPC 根据动力学、代价和约束在线求解控制序列。两者都能滚动执行，但获得动作的方法不同。')
p('自测 6  模型每次预测 10 步动作，是否就说明它在解 MPC 优化问题？')

page(7,'训练数据如何变成策略')
h('一条轨迹包含什么')
p('一条 episode 是一次完整任务或交互过程。它通常包含按时间对齐的图像、语言任务、机器人状态、动作、结束标志和相关元数据。某一步的训练样本可写成 (oₜ, l, aₜ)，动作块模型则使用 (oₜ, l, aₜ:ₜ₊H₋₁)。实际模型也可能输入多帧历史。')
h('行为克隆的基本目标')
code('预测动作 = policy(观察, 指令)\n损失 = distance(预测动作, 示教动作)\n更新模型参数，使训练损失下降')
p('以上是解释性伪代码。连续回归可采用平方误差，离散 token 可采用交叉熵，生成策略则采用对应的扩散或流匹配训练目标。监督目标应与动作编码一致，不能只看最终输出是数值就假定所有模型都使用同一种损失。')
h('训练流程中的关键步骤')
p('先明确动作接口，再采集或选择数据。按 episode 或更严格的场景划分训练、验证、测试集；只使用训练集计算需要学习的统计量；完成预处理与时间对齐；训练并保存检查点；最后在独立闭环任务中比较不同检查点。')
p('随机拆分相邻视频帧容易造成泄漏：同一次抓取的连续帧很相似，模型可能在测试时看到几乎相同的场景。若目标是新场景泛化，应按场景划分；若目标是新物体，应明确保留未见物体。')
h('离线误差与闭环偏移')
p('行为克隆只学习示教分布里的动作。机器人实际执行时，一点位置偏差就可能进入示教里没有覆盖的状态，后续错误逐渐积累。增加合适的状态覆盖、纠错示教或交互数据可能有帮助；具体是否改善，需要做对照实验。')
p('同一任务有多个正确动作。评价模型时不能只比较是否与某条示教逐点相等，也要看它是否完成目标、是否符合约束。')
p('自测 7  为什么训练动作误差很小，机器人仍可能抓不到物体？请分别从数据和闭环执行解释。')

page(8,'数据质量与跨机器人迁移')
h('大规模数据解决什么')
p('Open X-Embodiment 汇集多种机器人的数据，推动跨机器人训练与 RT-X 研究。[10] 数据多样性能帮助模型看到更广的物体、环境和动作，但不同数据源往往需要动作、状态和图像接口的映射。')
p('具身 embodiment 指机器人的身体形态、传感器、执行器与运动能力。六轴机械臂、双臂平台和平衡底盘并不拥有相同的动作空间。共享网络并不消除这些差异；通常需要匹配输入输出、选择兼容检查点，或用目标平台数据适配。')
h('先确认时间与几何一致')
p('图像时间戳与动作时间戳需要对应。若图像总是晚于动作，模型可能学到错位的因果关联。还要记录相机位置、内外参、关节零位、末端工具定义与采样周期。部署时改变镜头或控制频率，会影响已学到的映射。')
h('数据记录的最小清单')
table(['字段','记录内容'],[
('任务','语言指令 任务编号 成功或失败条件'),
('观察','图像 相机名称 时间戳 状态及其单位'),
('动作','动作向量 维度定义 坐标系 控制模式'),
('轨迹','episode 编号 起止时间 中止原因'),
('版本','机器人配置 标定 文件格式 预处理与代码版本')],[3.0,13.8])
h('示教覆盖与失败数据')
p('多次录制完全相同的位置，可能只增加重复性；改变目标位置、光照或摆放方式，才能检查你希望学习的变化。是否保留失败片段取决于训练目标：行为克隆需要合理监督，强化学习或失败检测可能需要结果标签。不能把所有错误动作都当成正确示教。')
p('仿真便于重复试验和生成数据；真实世界有接触、摩擦、遮挡和传感器噪声。sim-to-real 描述从仿真向真实机器人迁移的过程，需要评估这些差异。')
p('自测 8  给机械臂训练的 VLA 接到平衡底盘上，最先需要检查什么？')

page(9,'代表模型及其研究问题')
p('以下按技术路线选择代表方法，年份指首次公开研究的时间。不同论文的机器人、数据和任务不同，不据此给出统一名次。')
table(['工作','核心方向','初学者关注点'],[
('RT-1  2022','图像与指令条件的机器人动作 token','多任务策略与实时动作接口 [3]'),
('RT-2  2023','把动作编码为 token 与视觉语言数据共训','互联网语义如何迁移到控制 [4]'),
('OpenVLA  2024','开放的预训练 VLA 与微调流程','从预训练 VLM 到动作输出 [5]'),
('π0  2024','VLM 与流匹配动作生成结合','连续动作块与动作专家 [6]'),
('π0.5  2025','通过异构任务共训研究开放环境泛化','语言语义与动作学习的结合 [7]'),
('OpenVLA O FT  2025','研究微调中的动作块与解码效率','原始模型与后续改进的差别 [11]')],[3.5,6.1,7.2])
h('RT 系列给出的基础问题')
p('RT-1 表明 Transformer 可以用于多任务机器人控制，但它不等同于把一个通用聊天模型直接接上机器人。RT-2 更明确地把机器人动作纳入视觉语言模型的 token 输出，研究网络知识对语义任务的帮助。[3][4]')
h('OpenVLA 的开放研究价值')
p('原始 OpenVLA 是 7B 参数模型，使用约 97 万条机器人示范轨迹训练。它结合视觉编码器、投影模块与语言模型骨干，输出动作 token，提供权重和训练代码。[5] “开源可下载”意味着可以开展研究，不意味着任何机器人都能未经适配直接工作。')
h('阅读 π 系列时注意动作生成')
p('π0 用流匹配生成连续动作；π0.5 研究来自多机器人、高层语义和其他数据的共训以改善泛化。[6][7] openpi 提供相关实现与检查点；实现支持、模型权重和论文实验设置需要分别核对。[12]')
p('自测 9  为什么“参数更多”不够用来判断某模型更适合你的机器人？')

page(10,'如何判断 VLA 真的有效')
h('先把成功定义写清楚')
p('抓杯子的成功条件可以是：选对目标、稳定抬起、放入指定区域、松开夹爪并保持物体稳定。还应明确允许的时间、干预次数和重试方式。只报告“抓起来过”会掩盖选错对象或放置失败。')
p('成功率 = 成功次数 / 总试验次数。假设 20 次成功 16 次，观测成功率为 80%；样本仍较少，不能把这个值解释成所有环境下的固定概率。应同时记录测试条件、随机种子、总次数和失败分布。')
h('同时评价能力与执行成本')
table(['指标','回答的问题'],[
('任务成功率','是否完成定义好的任务'),
('语言遵循','更换指令后 是否改变目标或行为'),
('泛化','新物体 新摆放 新背景或新机器人上怎样'),
('延迟与吞吐','端到端响应多久 动作更新能否跟上'),
('物理约束','碰撞 限位 超速或接触异常是否出现'),
('恢复能力','偏差或干扰发生后 能否继续完成任务')],[3.6,13.2])
h('延迟要从相机算到执行器')
p('测量应包含采图、预处理、网络传输、模型推理、动作解码和发送。平均延迟会掩盖偶发长等待，还应查看高分位延迟和截止时间错过次数。若使用远端 GPU，通信故障和动作陈旧问题也属于闭环性能。')
h('用对照实验排除误解')
p('保持硬件、任务和测试条件一致，再比较模型或组件。可以改变指令、移除状态输入、换相机背景，观察具体影响。实验室内达到高成功率，不自动说明跨家庭、跨机器人或动态场景都有效。')
p('LIBERO 提供语言条件的机器人学习评测环境，可作为仿真入门入口。[13] 读论文时需要核对任务集、数据划分和评测协议；视频展示常经过加速或筛选，不能替代完整统计。')
p('自测 10  两个模型分别在不同机器人上报告 90% 和 80%，能否断言前者更好？')

page(11,'VLA 与你的平衡底盘和 MPC')
diagram('mpc_layers','layers')
h('分工取决于动作接口')
p('VLA 可以直接输出某种机器人动作，也可以在工程系统中提供上层目标或参考。对平衡底盘，一个便于理解的方案是让视觉语言系统确定目标与任务，让规划层给出速度和转向参考，再由状态估计和 MPC 维持平衡并跟踪参考。这是设计建议，并不是每篇 VLA 论文都采用的架构。')
p('如果视觉语言模块只输出目标文字或目标点，通常更准确地称为视觉语言规划系统；动作仍由独立控制器生成。若模型学习视觉与语言到机器人动作的映射，则符合常见 VLA 用法。命名应反映实际输入输出。')
h('MPC 是怎样获得控制量的')
p('MPC 用动力学模型预测未来状态，并寻找满足约束、代价较低的控制序列。以教学示意为例：xₜ₊₁ = A xₜ + B uₜ，代价包含状态偏离参考的误差和控制量大小。每次执行首个控制量，再测量状态并重新求解。')
p('对底盘，x 可以包含倾角、角速度、位置和速度；u 可以是轮端力矩或适配后的控制输入。真实状态顺序、模型参数、约束及采样周期必须根据你的硬件定义，不能照搬机械臂或倒立摆例子的矩阵。')
h('两种系统怎样衔接')
p('上层参考要经过单位转换、速度限制与变化率限制，再进入底层控制。底层反馈不应等待一次大型模型推理才更新。若参考过期或上层暂时无响应，需要预先定义降速、保持平衡或停止任务的行为。具体控制周期应通过动力学与算力测试确定。')
p('你已经跑通 TinyMPC，说明示例编译与运行链路成立。下一步是理解模型、状态、代价和约束，再在仿真中验证自己的底盘。已有机械臂 VLA 检查点能否适配底盘，还需要目标动作数据与实验。[14]')
p('自测 11  大模型暂时卡住时，维持平衡的反馈控制为什么还应继续运行？')

page(12,'分阶段实践路线')
p('下面按每周约 5 至 8 小时设计，是可调整的学习安排。先做到每阶段的验收目标，再决定是否增加算力与硬件投入。')
table(['阶段','学习与实验','验收结果'],[
('第 1 周','Python 数组 文件路径 状态动作与坐标系','能解释一条观察和动作的每个字段'),
('第 2 周','神经网络训练与推理 简单行为克隆','能区分训练损失与测试表现'),
('第 3 周','阅读 RT-2 和 OpenVLA 架构 数据样本','画出输入到执行器的完整路径'),
('第 4 周','官方仿真或离线推理示例','记录输出形状 单位 检查点与延迟'),
('第 5 周','改变指令或摆放 做固定协议评测','生成成功率与失败类型记录'),
('第 6 周','底盘仿真 MPC 参考接口 语言目标','比较无语言基线与加入上层后的行为')],[2.2,8.6,6.0])
h('普通电脑也能先完成的任务')
p('阅读一条数据、画坐标系、验证动作反归一化、写评测记录和运行简单数值实验，都不需要先训练大模型。你当前的 WSL Ubuntu 可用于学习代码和控制仿真，但安装 WSL 并不自动获得可用的 GPU、匹配的 CUDA 或足够显存。')
p('开始运行大模型前，应查看对应仓库对操作系统、GPU、内存和软件版本的要求。推理、LoRA 微调和全参数训练的资源需求不同，无法只根据参数数量精确判断。本讲义不假定你已具备训练设备。[12][15]')
h('从两个官方入口选择实践')
p('OpenVLA 适合学习从 VLM 到动作 token 及微调；openpi 适合学习连续动作块与动作服务；LeRobot 提供数据、训练和机器人学习工作流。[5][12][15] 先选一个项目并固定代码版本，不要把多个项目的依赖安装进同一个环境。')
p('自测 12  没有大显存 GPU，仍能完成哪些有价值的 VLA 学习任务？')

page(13,'一个不需要机器人硬件的小练习')
h('检查动作反归一化与限幅')
p('这个练习只说明数据接口，使用 Python 标准库即可运行。它不是 VLA 模型，也不连接电机。我们假设模型给出标准化的二维位移增量，每个维度按 a=z×σ+μ 还原，再限制为 ±0.01 米。')
code('z = [0.5, -1.2]\nmu = [0.0, 0.0]\nsigma = [0.02, 0.02]  # 每一维的单位为米\nlimit = 0.01\n\nraw = [v * s + m for v, s, m in zip(z, sigma, mu)]\nsafe = [max(-limit, min(limit, v)) for v in raw]\n\nprint("raw:", raw)\nprint("limited:", safe)')
p('期望结果：原始增量为 [0.01, -0.024] 米；限幅后约为 [0.01, -0.01] 米。浮点数显示可能有微小差异。限幅只限制单步数值，不能保证轨迹无碰撞，也不能保证平衡稳定。')
h('把每一步的约定写在纸上')
p('输入图像：有几台相机，宽高是多少，RGB 还是 BGR？状态：各维度是什么，单位是什么？输出：单步还是动作块，绝对值还是增量？时间：动作间隔是多少，推理多久更新？执行：哪个控制器消费这些输出，过期时怎么办？')
h('闭环的解释性伪代码')
code('while task_active:\n    observation = read_sensors_with_timestamp()\n    chunk = policy(observation, instruction)\n    actions = decode_with_matching_statistics(chunk)\n    actions = check_units_frames_limits(actions)\n    execute_first_k_actions(actions)\n    log_observation_action_result()')
p('真实系统还要处理异步时序、异常、结束判定与独立运行的底层控制。这里用顺序伪代码帮助理解数据流，不能直接作为平衡底盘控制程序。')
p('自测 13  如果把 sigma 改成 0.2，同样的模型输出会怎样变化？你能解释为什么统计量是模型部署接口的一部分吗？')

page(14,'常见误解与术语表')
h('学习中容易混淆的地方')
p('“能聊天就能控制机器人”：语言输出仍需要与机器人动作相连。“有图像就是 VLA”：还需看语言是否参与，以及是否输出动作。“端到端无需工程”：时间同步、单位与执行器接口仍是系统的一部分。“仿真成功就能上实机”：接触和传感器差异仍需验证。')
p('“看一段视频就等于学会”：普通视频往往没有准确动作标签，如何利用视频属于具体研究方法。“动作越平滑越成功”：平滑性与正确目标、约束满足、成功率是不同指标。“换一个中文指令即可迁移”：语言覆盖与机器人动作覆盖需要分别检查。')
table(['术语','中文与实际含义'],[
('Observation','观察 传感器提供的输入'),
('State','状态 描述机器人与环境的变量'),
('Policy','策略 从观察或状态到动作的映射'),
('Action space','动作空间 允许输出的动作形式与范围'),
('Embodiment','具身 机器人的身体 传感器与动作能力'),
('Grounding','落地关联 把语言概念对应到对象与行为'),
('Episode / Rollout','一段任务轨迹 / 策略执行的交互过程'),
('Action chunk','动作块 一次预测的多个时间步动作'),
('Checkpoint','检查点 保存的模型参数及相关配置'),
('Fine tuning / LoRA','微调 / 使用低秩更新的参数高效微调'),
('Zero shot','零样本 相对指定任务或设置没有额外适配'),
('OOD','分布外 相对训练分布存在某种变化'),
('Latency','延迟 从请求或采集到结果或执行的时间'),
('Closed loop','闭环 利用新反馈更新后续行为')],[4.4,12.4])
p('“零样本”和“分布外”都要说明针对什么变化：新物体、新任务、新背景或新机器人，不是任意场景通用的同义词。')

page(15,'自测答案与下一步阅读')
p('1  视觉目标不能直接给出可靠的快速姿态反馈；IMU 用于估计倾角与角速度等平衡相关量。\n2  是推理。改指令可影响条件输入，但不能补上未知动作接口与目标机器人经验。\n3  还需动作生成、标定与坐标变换、解码、控制器和硬件执行。\n4  覆盖 1 秒；每执行 4 步是 0.2 秒，约 5 次更新每秒，真实周期还受延迟影响。\n5  不能。查归一化、单位、坐标系、绝对或增量定义及采样周期。')
p('6  不是。预测动作块不代表使用动力学和约束在线优化。\n7  测试可能超出示教分布；时延、错位输入或小误差累积也会导致闭环失败。\n8  查动作维度、单位、坐标系、输入状态和检查点覆盖的具身。\n9  还要考虑数据、动作空间、任务、延迟、硬件和微调条件。\n10  不能。需在匹配任务与协议下比较，并报告次数与不确定性。')
p('11  平衡是持续的反馈任务，上层延迟不应中断姿态稳定控制。\n12  可检查数据、学习接口、练习反归一化、读论文、写评测协议和做控制仿真。\n13  原始增量变为 [0.1, -0.24] 米，限幅结果仍为 [0.01, -0.01]；限幅掩盖了尺度错误，所以要同时检查原始值。')
h('读论文时固定问六个问题')
p('输入是什么？输出的动作怎样定义？训练数据来自哪里？测试改变了什么条件？成功怎样统计？部署需要哪些接口和计算资源？先在摘要、架构图和实验设置里寻找这些答案，再读训练目标。')
h('你的第一个小项目')
p('先完成“语言指定目标，底盘在仿真中移动并保持平衡”。固定两三个简单目标，使用明确的参考接口和 MPC 基线，记录成功条件、参考速度、姿态误差与控制输入。可以先用规则把指令变成目标，验证控制链路，再研究是否需要学习式 VLA。这个阶段的目标是建立可测量的完整闭环。')
p('若以后尝试 VLA，比较同一系统中加入模型前后的目标选择、参考生成与执行结果。若主要困难是姿态稳定或轮端模型，先解决控制问题；若主要困难是语言条件下识别目标与行为迁移，再投入相应的模型与数据工作。')

page(16,'一手资料与阅读顺序')
p('资料核对日期为 2026 年 10 月 14 日。模型、依赖和仓库会更新；安装时以对应官方版本说明为准。正文按编号引用，以下链接可直接打开。')
refs=[
('1','Attention Is All You Need','2017  Transformer 的原始论文','https://arxiv.org/abs/1706.03762'),
('2','LoRA Low Rank Adaptation of Large Language Models','2021  参数高效微调的原始论文','https://arxiv.org/abs/2106.09685'),
('3','RT 1 Robotics Transformer for Real World Control at Scale','2022  多任务机器人动作策略','https://arxiv.org/abs/2212.06817'),
('4','RT 2 Vision Language Action Models','2023  语义知识到机器人控制 项目页及论文','https://robotics-transformer2.github.io/'),
('5','OpenVLA An Open Source Vision Language Action Model','2024  模型结构 数据 权重与代码入口','https://openvla.github.io/'),
('6','π0 A Vision Language Action Flow Model','2024  连续动作与流匹配','https://arxiv.org/abs/2410.24164'),
('7','π0.5 A VLA with Open World Generalization','2025  异构数据共训与泛化','https://arxiv.org/abs/2504.16054'),
('8','Diffusion Policy','2023  条件动作生成的基础路线','https://diffusion-policy.cs.columbia.edu/'),
('9','Flow Matching for Generative Modeling','2022  流匹配基础','https://arxiv.org/abs/2210.02747'),
('10','Open X Embodiment','2023  跨机器人数据与 RT X 项目','https://robotics-transformer-x.github.io/'),
('11','Fine Tuning VLA Models Optimizing Speed and Success','2025  OpenVLA O FT 后续改进','https://arxiv.org/abs/2502.19645'),
('12','openpi 官方代码','π 系列实现 检查点与部署要求','https://github.com/Physical-Intelligence/openpi'),
('13','LIBERO 官方代码','仿真基准 任务与评测入口','https://github.com/Lifelong-Robot-Learning/LIBERO'),
('14','TinyMPC 官方代码','MPC 求解器 示例与代码生成','https://github.com/TinyMPC/TinyMPC'),
('15','LeRobot 官方文档','机器人数据 训练与实践工作流','https://huggingface.co/docs/lerobot/index')]
def hyperlink(para,label,url):
    rel=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'23577D');pr.append(col)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');pr.append(size)
    r.append(pr);txt=OxmlElement('w:t');txt.text=label;r.append(txt);link.append(r);para._p.append(link)
for no,title,description,url in refs:
    para=doc.add_paragraph();para.paragraph_format.space_after=Pt(5)
    para.paragraph_format.line_spacing=Pt(13)
    para.add_run(f'[{no}] ').font.size=Pt(9);hyperlink(para,title,url)
    para.add_run('\n'+description).font.size=Pt(9)
p('建议顺序  先读 [4][5] 的图与项目说明，再读 [8][6][7]，实践时查 [12][15]。需要推导时回到 [1][2][9]。结合平衡底盘学习控制时，同时阅读 [14] 并检查自己的动力学模型。')
dest=OUT/'VLA机器人模型入门讲义.docx'
doc.save(dest)
print(dest)
print('Paragraphs:',len(doc.paragraphs),'Tables:',len(doc.tables))
